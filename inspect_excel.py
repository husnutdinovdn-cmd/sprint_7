# -*- coding: utf-8 -*-
import openpyxl
import sys
sys.stdout.reconfigure(encoding='utf-8')
path = r'D:\Иван Иванов — диплом_Инженер по тестированию-v-2-4-25.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
for name in ['Задание 1 чек-лист', 'Задание 1 данные валидации', 'Баг-репорты', 'Задание 1 баги вне тестовой док']:
    ws = wb[name]
    mr, mc = ws.max_row or 0, ws.max_column or 0
    print('===', name, 'rows:', mr, 'cols:', mc)
    for row in range(1, min(6, mr + 1)):
        r = [ws.cell(row=row, column=c).value for c in range(1, min(10, mc + 1))]
        print(r)
wb.close()
