# -*- coding: utf-8 -*-
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document

for path in [
    r'C:\Users\222\Desktop\диплом Задание 1.docx',
    r'C:\Users\222\Desktop\диплом задание 2.docx',
]:
    print('='*60, path, '='*60)
    try:
        doc = Document(path)
        for p in doc.paragraphs:
            print(p.text)
        for t in doc.tables:
            for row in t.rows:
                print(' | '.join(cell.text.strip() for cell in row.cells))
    except Exception as e:
        print('Error:', e)
    print()
