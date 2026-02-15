# -*- coding: utf-8 -*-
import openpyxl
import sys
sys.stdout.reconfigure(encoding='utf-8')
path = r'D:\Иван Иванов — диплом_Инженер по тестированию-v-2-4-25.xlsx'
wb = openpyxl.load_workbook(path, data_only=True)
name = 'Задание 2 тест-кейсы'
ws = wb[name]
mr, mc = ws.max_row or 0, ws.max_column or 0
print('Sheet:', name, 'rows:', mr, 'cols:', mc)
for r in range(1, min(35, mr + 1)):
    row = [ws.cell(row=r, column=c).value for c in range(1, min(8, mc + 1))]
    print(r, row[:7])
wb.close()
